"""
EduPPT Creator - PowerPoint Generator
Version: 1.0

Tạo PowerPoint (.pptx) từ dữ liệu slide dạng JSON.
Sử dụng thư viện python-pptx.
"""

import argparse
import json
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


# ============================================================
# CẤU HÌNH
# ============================================================

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

FONT_TITLE = "Arial"
FONT_BODY = "Arial"

TITLE_SIZE = Pt(28)
BODY_SIZE = Pt(20)
SMALL_SIZE = Pt(14)


# ============================================================
# HÀM TẠO TEXT
# ============================================================

def add_textbox(slide, text, left, top, width, height,
                font_size=BODY_SIZE,
                bold=False,
                align=PP_ALIGN.LEFT):

    textbox = slide.shapes.add_textbox(
        left,
        top,
        width,
        height
    )

    text_frame = textbox.text_frame
    text_frame.clear()

    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align

    run = paragraph.add_run()
    run.text = str(text)

    run.font.name = FONT_BODY
    run.font.size = font_size
    run.font.bold = bold

    return textbox


# ============================================================
# TẠO TIÊU ĐỀ
# ============================================================

def add_title(slide, title):

    add_textbox(
        slide,
        title,
        Inches(0.7),
        Inches(0.35),
        Inches(11.9),
        Inches(0.8),
        font_size=TITLE_SIZE,
        bold=True
    )


# ============================================================
# TẠO NỘI DUNG BULLET
# ============================================================

def add_bullets(slide, bullets):

    textbox = slide.shapes.add_textbox(
        Inches(0.9),
        Inches(1.45),
        Inches(11.5),
        Inches(5.2)
    )

    text_frame = textbox.text_frame
    text_frame.clear()

    for index, bullet in enumerate(bullets):

        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()

        paragraph.text = str(bullet)
        paragraph.font.name = FONT_BODY
        paragraph.font.size = BODY_SIZE

        paragraph.level = 0

    return textbox


# ============================================================
# TẠO SLIDE TIÊU ĐỀ
# ============================================================

def create_title_slide(prs, slide_data):

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    title = slide_data.get(
        "title",
        "EduPPT Creator"
    )

    subtitle = slide_data.get(
        "subtitle",
        ""
    )

    add_textbox(
        slide,
        title,
        Inches(1.0),
        Inches(2.2),
        Inches(11.3),
        Inches(1.2),
        font_size=Pt(34),
        bold=True,
        align=PP_ALIGN.CENTER
    )

    if subtitle:

        add_textbox(
            slide,
            subtitle,
            Inches(1.5),
            Inches(3.6),
            Inches(10.3),
            Inches(1.0),
            font_size=Pt(22),
            align=PP_ALIGN.CENTER
        )

    return slide


# ============================================================
# TẠO SLIDE NỘI DUNG
# ============================================================

def create_content_slide(prs, slide_data):

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    title = slide_data.get(
        "title",
        "Nội dung"
    )

    bullets = slide_data.get(
        "bullets",
        []
    )

    add_title(slide, title)

    if bullets:
        add_bullets(slide, bullets)

    return slide


# ============================================================
# TẠO SLIDE
# ============================================================

def create_slide(prs, slide_data):

    slide_type = slide_data.get(
        "type",
        "content"
    )

    if slide_type == "title":

        return create_title_slide(
            prs,
            slide_data
        )

    return create_content_slide(
        prs,
        slide_data
    )


# ============================================================
# ĐỌC JSON
# ============================================================

def load_json(input_file):

    with open(
        input_file,
        "r",
        encoding="utf-8"
    ) as file:

        return json.load(file)


# ============================================================
# TẠO POWERPOINT
# ============================================================

def create_presentation(data, output_file):

    prs = Presentation()

    # Thiết lập khổ 16:9
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slides = data.get(
        "slides",
        []
    )

    if not slides:

        raise ValueError(
            "Không tìm thấy danh sách slides trong dữ liệu."
        )

    for slide_data in slides:

        create_slide(
            prs,
            slide_data
        )

    # Tạo thư mục nếu chưa tồn tại
    output_dir = os.path.dirname(
        output_file
    )

    if output_dir:
        os.makedirs(
            output_dir,
            exist_ok=True
        )

    prs.save(
        output_file
    )

    return output_file


# ============================================================
# MAIN
# ============================================================

def main():

    parser = argparse.ArgumentParser(
        description="EduPPT Creator - Tạo PowerPoint từ JSON"
    )

    parser.add_argument(
        "--input",
        required=True,
        help="Đường dẫn file JSON chứa nội dung slide"
    )

    parser.add_argument(
        "--output",
        required=True,
        help="Đường dẫn file PowerPoint đầu ra"
    )

    args = parser.parse_args()

    print("======================================")
    print("EduPPT Creator")
    print("PowerPoint Generator")
    print("======================================")

    print(
        f"Đang đọc dữ liệu: {args.input}"
    )

    data = load_json(
        args.input
    )

    print(
        f"Tìm thấy {len(data.get('slides', []))} slide."
    )

    print(
        "Đang tạo PowerPoint..."
    )

    output = create_presentation(
        data,
        args.output
    )

    print(
        "Đã tạo PowerPoint thành công!"
    )

    print(
        f"File: {output}"
    )


if __name__ == "__main__":

    main()
